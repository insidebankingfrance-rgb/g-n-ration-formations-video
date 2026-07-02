"""Extrait le texte de chaque slide d'un support PowerPoint (.pptx)
vers un fichier JSON structuré, utilisable pour écrire le script de narration.
"""

import argparse
import json
from pathlib import Path

from pptx import Presentation


def extract_slide(slide, index):
    title = None
    bullets = []

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue

        is_title = shape == slide.shapes.title
        if is_title and title is None:
            title = text
        else:
            bullets.extend([line.strip() for line in text.split("\n") if line.strip()])

    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()

    return {
        "index": index,
        "title": title or f"Slide {index + 1}",
        "bullets": bullets,
        "notes": notes,
    }


def extract(pptx_path):
    prs = Presentation(pptx_path)
    return [extract_slide(slide, i) for i, slide in enumerate(prs.slides)]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", help="Chemin vers le fichier .pptx source")
    parser.add_argument(
        "-o", "--output", default="data/slides.json", help="Fichier JSON de sortie"
    )
    args = parser.parse_args()

    slides = extract(args.pptx)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"{len(slides)} slides extraites -> {out_path}")


if __name__ == "__main__":
    main()
